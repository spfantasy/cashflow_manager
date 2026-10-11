"""Web 页面测试。"""

from __future__ import annotations

import io
import os
import re

from openpyxl import load_workbook

from .conftest import INVOICE_FILE, ORDERS_FILE, TRANSACTIONS_FILE, samples_available


def test_all_pages_render(client):
    for path in ("/", "/details", "/imports", "/listings", "/manual", "/settings"):
        response = client.get(path)
        assert response.status_code == 200, path
        assert "现金流" in response.get_data(as_text=True)


def test_settings_seed_and_protected_unknown(client):
    page = client.get("/settings").get_data(as_text=True)
    assert "TEST_USER" in page
    assert "Unknown" in page
    # 初始正则就是页面6 可编辑的文本
    assert "^.* (TEST_USER)$" in page


def test_settings_save_and_rematch(app, client):
    response = client.post(
        "/settings/seller/save",
        data={
            "seller_name": "李楚吟",
            "match_regex": r"^.*Lot\(3\)$",
            "sort_order": "10",
            "share_membership": "1",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "李楚吟" in response.get_data(as_text=True)

    from cashflow.db import connect
    from cashflow.repository import participants

    connection = connect(app.config["DB_PATH"])
    try:
        assert "李楚吟" in participants(connection)
    finally:
        connection.close()


def test_settings_rejects_invalid_regex(client):
    response = client.post(
        "/settings/seller/save",
        data={"seller_name": "bad", "match_regex": "([", "sort_order": "10"},
        follow_redirects=True,
    )
    assert "正则表达式无效" in response.get_data(as_text=True)


def test_manual_cashflow_creation(client):
    response = client.post(
        "/manual",
        data={
            "item_id": "168063599242",
            "amount": "-21.95",
            "cashflow_date": "2026-01-16",
            "reminder": "退款编号 5180906651",
        },
        follow_redirects=True,
    )
    body = response.get_data(as_text=True)
    assert "已创建手工帐" in body
    assert "168063599242-" in body

    details = client.get("/details?item_id=168063599242").get_data(as_text=True)
    assert "手工帐" in details
    assert "-21.95" in details


def test_manual_rejects_bad_amount(client):
    response = client.post(
        "/manual",
        data={"item_id": "x", "amount": "abc", "cashflow_date": "2026-01-16"},
        follow_redirects=True,
    )
    assert "金额必须填写且为数字" in response.get_data(as_text=True)


def test_listing_adjustment_updates_seller(app, client):
    from cashflow.db import connect
    from cashflow.repository import insert_item

    connection = connect(app.config["DB_PATH"])
    try:
        insert_item(
            connection,
            {"item_id": "168000000001", "item_name": "Test Title", "seller_name": "Unknown"},
        )
        connection.commit()
    finally:
        connection.close()

    response = client.post(
        "/listings",
        data={"item_id": "168000000001", "seller_name": "李楚吟"},
        follow_redirects=True,
    )
    assert "已调整为「李楚吟」" in response.get_data(as_text=True)

    connection = connect(app.config["DB_PATH"])
    try:
        row = connection.execute(
            "SELECT seller_name, seller_manual FROM ebay_item WHERE item_id='168000000001'"
        ).fetchone()
        assert row["seller_name"] == "李楚吟"
        assert row["seller_manual"] == 1
    finally:
        connection.close()


def test_flows_page_filters_by_user(app, client):
    """账户级现金流的 item_id 是 ebay账户-用户名，要能用用户名筛出来。"""
    client.post(
        "/settings/seller/save",
        data={"seller_name": "王五", "match_regex": "", "sort_order": "20", "share_membership": "1"},
    )
    client.post(
        "/manual",
        data={"item_id": "ebay账户-王五", "amount": "-1.97", "cashflow_date": "2026-01-17"},
    )
    body = client.get("/flows?seller=王五").get_data(as_text=True)
    assert "-1.97" in body
    assert "共 1 笔现金流" in body

    body = client.get("/flows?seller=没有这个人").get_data(as_text=True)
    assert "共 0 笔现金流" in body


@samples_available
def test_export_flows_to_xlsx(app, client):
    from cashflow.db import connect
    from cashflow.importer import ORDERS_KIND, import_file

    connection = connect(app.config["DB_PATH"])
    try:
        import_file(connection, ORDERS_KIND, ORDERS_FILE)
    finally:
        connection.close()

    response = client.get("/flows/export?item_id=")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook.active
    assert sheet["A1"].value == "现金流ID"
    assert sheet.max_row > 600  # 311 个 item × 2 笔 + 表头 + 合计行


@samples_available
def test_import_page_uploads_files(app, client):
    data = {
        "orders_files": (open(ORDERS_FILE, "rb"), os.path.basename(ORDERS_FILE)),
        "transactions_files": (open(TRANSACTIONS_FILE, "rb"), os.path.basename(TRANSACTIONS_FILE)),
        "invoice_files": (open(INVOICE_FILE, "rb"), os.path.basename(INVOICE_FILE)),
    }
    response = client.post(
        "/imports", data=data, content_type="multipart/form-data", follow_redirects=True
    )
    assert response.status_code == 200
    body = response.get_data(as_text=True)
    assert "新增物品" in body
    assert "311" in body

    # 有需要人工处理的记录时会弹出对话框（数据通过 JSON 传给前端）
    assert "issues-data" in body
    assert "需要人工处理" in body


def test_import_page_rejects_unknown_extension(client):
    response = client.post(
        "/imports",
        data={"orders_files": (io.BytesIO(b"x"), "bad.txt")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "不支持的文件类型" in response.get_data(as_text=True)


def test_details_delete_removes_rows(client):
    client.post(
        "/manual",
        data={"item_id": "168099999999", "amount": "1.00", "cashflow_date": "2026-01-17"},
    )
    details = client.get("/details?item_id=168099999999").get_data(as_text=True)
    match = re.search(r'name="cashflow_id" value="([^"]+)"', details)
    assert match, "详情页应该有可勾选删除的现金流"
    cashflow_id = match.group(1)

    response = client.post(
        "/details/delete?item_id=168099999999",
        data={"cashflow_id": cashflow_id},
        follow_redirects=True,
    )
    assert "已删除 1 行现金流" in response.get_data(as_text=True)
    assert "没有符合条件的现金流" in response.get_data(as_text=True)


def test_details_filters_are_or_relation(app, client):
    """item_id 与 order_id 是「或」关系：任意一个匹配就会列出。"""
    from cashflow.db import connect
    from cashflow.repository import insert_cashflow, insert_item

    connection = connect(app.config["DB_PATH"])
    try:
        for item_id, order_id in (("A-1", "ORDER-A"), ("B-2", "ORDER-B")):
            insert_item(connection, {"item_id": item_id, "order_id": order_id})
            insert_cashflow(
                connection,
                {
                    "cashflow_id": f"MANUAL-{item_id}",
                    "item_id": item_id,
                    "order_id": order_id,
                    "cashflow_source": "手工帐",
                    "cashflow_type": "手工帐",
                    "amount": "1.00",
                    "cashflow_date": 20260117,
                },
            )
        connection.commit()
    finally:
        connection.close()

    # 只给 item_id，只应命中 A-1
    body = client.get("/details?item_id=A-1").get_data(as_text=True)
    assert "A-1" in body and "B-2" not in body

    # item_id 命中 A-1、order_id 命中 ORDER-B，或关系应同时列出两条
    body = client.get("/details?item_id=A-1&order_id=ORDER-B").get_data(as_text=True)
    assert "A-1" in body and "B-2" in body
