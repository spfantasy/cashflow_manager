"""导入器公共设施：读文件、定位表头、取列、收集导入结果。"""

from __future__ import annotations

import collections
import csv
import io
import os
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence

from ..money import cell_is_empty, cell_str, compact

CSV_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5", "cp1252")


# --------------------------------------------------------------------------- #
# 读取文件
# --------------------------------------------------------------------------- #
@dataclass
class Table:
    header: List[str]
    rows: List[List]

    def index_of(self, *candidates: str) -> int:
        """按「去空白后包含」的方式找列，返回列下标，找不到返回 -1。"""
        normalized = [compact(name) for name in self.header]
        for candidate in candidates:
            key = compact(candidate)
            if not key:
                continue
            if key in normalized:
                return normalized.index(key)
        for candidate in candidates:
            key = compact(candidate)
            if not key:
                continue
            for position, name in enumerate(normalized):
                if key and key in name:
                    return position
        return -1


def _read_csv(path: str) -> List[List[str]]:
    last_error: Optional[Exception] = None
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                content = handle.read()
            break
        except UnicodeDecodeError as exc:  # 换下一种编码
            last_error = exc
    else:
        raise ValueError(f"无法识别文件编码：{last_error}")
    return [list(row) for row in csv.reader(io.StringIO(content))]


def _read_xlsx(path: str) -> List[List]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("读取 xlsx 需要安装 openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_file(path: str) -> List[List]:
    extension = os.path.splitext(path)[1].lower()
    if extension in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if extension == ".xls":
        raise ValueError("不支持旧版 .xls，请用 Excel 另存为 .xlsx 或 .csv")
    return _read_csv(path)


def load_table(path: str, header_keys: Sequence[str]) -> Table:
    """读文件并定位表头行（表头前可能有若干说明行）。"""
    raw = read_file(path)
    header_row = -1
    for position, row in enumerate(raw):
        joined = [compact(cell) for cell in row]
        if all(any(key == name or key in name for name in joined) for key in header_keys):
            header_row = position
            break
    if header_row < 0:
        raise ValueError(
            "没有找到预期的表头（缺少列：" + "、".join(header_keys) + "），请确认文件类型是否正确"
        )
    header = [cell_str(cell) for cell in raw[header_row]]
    candidates = [list(row) for row in raw[header_row + 1:] if any(cell_str(cell) for cell in row)]
    # 文件尾部常有统计行（如 "62,笔已下载记录"、"卖家用户名 : lcy-228"），
    # 它们比正常数据行短得多。以「出现次数最多的行长」为准，丢掉明显偏短的行，
    # 这样既不会把统计行当数据，也不会误伤列数比表头少的文件。
    threshold = 2
    if candidates:
        typical = collections.Counter(len(row) for row in candidates).most_common(1)[0][0]
        threshold = max(2, min(typical, len(header)) // 2)
    body = [row for row in candidates if len(row) >= threshold]
    return Table(header=header, rows=body)


# --------------------------------------------------------------------------- #
# 取列工具
# --------------------------------------------------------------------------- #
class Columns:
    """把「别名 -> 列下标」的解析集中起来，缺列时给出明确报错。"""

    def __init__(self, table: Table, aliases: Dict[str, Sequence[str]], required: Iterable[str] = ()):
        self.table = table
        self.map: Dict[str, int] = {}
        for name, candidates in aliases.items():
            self.map[name] = table.index_of(*candidates)
        missing = [name for name in required if self.map.get(name, -1) < 0]
        if missing:
            raise ValueError("文件缺少必要的列：" + "、".join(missing))

    def raw(self, row: Sequence, name: str):
        position = self.map.get(name, -1)
        if position < 0 or position >= len(row):
            return ""
        value = row[position]
        return "" if value is None else value

    def text(self, row: Sequence, name: str) -> str:
        return cell_str(self.raw(row, name))

    def present(self, row: Sequence, name: str) -> bool:
        """该列是否有真实取值（'--'、空串都算没有）。"""
        return not cell_is_empty(self.raw(row, name))


# --------------------------------------------------------------------------- #
# 导入结果
# --------------------------------------------------------------------------- #
@dataclass
class ImportIssue:
    """需要人工处理的失败项（页面3 会弹对话框提示）。"""

    reason: str
    reference: str = ""
    row_number: int = 0
    detail: str = ""

    def as_dict(self) -> dict:
        return {
            "reason": self.reason,
            "reference": self.reference,
            "row_number": self.row_number,
            "detail": self.detail,
        }


@dataclass
class ImportResult:
    file_name: str
    file_kind: str
    inserted_items: int = 0
    skipped_items: int = 0
    inserted_cashflows: int = 0
    skipped_cashflows: int = 0
    issues: List[ImportIssue] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def failed(self) -> int:
        return len(self.issues)

    def add_issue(self, reason: str, reference: str = "", row_number: int = 0, detail: str = "") -> None:
        self.issues.append(
            ImportIssue(reason=reason, reference=reference, row_number=row_number, detail=detail)
        )

    def note(self, message: str) -> None:
        if message not in self.notes:
            self.notes.append(message)

    def merge(self, other: "ImportResult") -> None:
        """把另一个文件的结果并入本文件结果（页面3 每个拖入框可能一次多个文件）。"""
        self.inserted_items += other.inserted_items
        self.skipped_items += other.skipped_items
        self.inserted_cashflows += other.inserted_cashflows
        self.skipped_cashflows += other.skipped_cashflows
        self.issues.extend(other.issues)
        for note in other.notes:
            self.note(note)
