"""Excel 导出。

说明：xlsx 内部就是 IEEE754 双精度，没有真正的十进制类型；导出时把金额以
2 位小数格式写入（数字格式 '0.00'），数据库里仍保留精确十进制字符串。
"""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .money import fmt_date, money

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)

FLOW_HEADERS = [
    ("现金流ID", "cashflow_id", 30),
    ("物品编号", "item_id", 18),
    ("订单编号", "order_id", 20),
    ("物品标题", "item_name", 52),
    ("用户", "seller_name", 14),
    ("买家", "buyer_username", 18),
    ("来源", "cashflow_source", 16),
    ("类型", "cashflow_type", 16),
    ("金额", "amount", 12),
    ("日期", "cashflow_date", 12),
    ("备注", "reminder", 46),
    ("追踪号码", "track_id", 30),
    ("lot", "lot", 6),
]


def _write_sheet(sheet, headers, rows: Sequence) -> None:
    for column, (title, _key, width) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        for column, (_title, key, _width) in enumerate(headers, start=1):
            if key == "amount":
                value: object = float(money(row["amount"]))
                cell = sheet.cell(row=row_index, column=column, value=value)
                cell.number_format = "0.00"
                continue
            if key == "cashflow_date":
                value = fmt_date(row["cashflow_date"])
            elif key == "lot":
                value = row["lot"] if "lot" in row.keys() else ""
            else:
                value = row[key] if key in row.keys() else ""
            sheet.cell(row=row_index, column=column, value=value)

    total = sum((money(row["amount"]) for row in rows), Decimal("0"))
    total_row = len(rows) + 2
    label_cell = sheet.cell(row=total_row, column=8, value="合计")
    label_cell.font = Font(bold=True)
    amount_cell = sheet.cell(row=total_row, column=9, value=float(total))
    amount_cell.number_format = "0.00"
    amount_cell.font = Font(bold=True)


def build_flow_workbook(rows: Sequence, title: str = "现金流") -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "现金流"
    _write_sheet(sheet, FLOW_HEADERS, rows)
    return workbook


def workbook_bytes(workbook: Workbook) -> io.BytesIO:
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def export_filename(prefix: str) -> str:
    return f"{prefix}-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
